"""Sheet + config -> Catalogue, with the validation rules of BRD section 4.4."""
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import yaml

from . import dims as dimslib
from .catalogue import Addon, Catalogue, Item, Product, Tier

log = logging.getLogger("printevr.loader")

MASTER_TAB = "Master Price List"
FLAGS_TAB = "Review Flags"
SAMPLE_TIER = "sample cost"
MASTER_COLUMNS = {
    "category": "Category",
    "product": "Product",
    "size": "Size / Spec",
    "option_1": "Option 1",
    "option_2": "Option 2",
    "tier_label": "Qty tier (as printed)",
    "qty_from": "Qty from",
    "unit": "Unit",
    "price": "Price ₹ (excl. GST)",
    "production_time": "Production time",
    "flag": "Flag",
}


class LoaderError(Exception):
    """The sheet or config is not safe to quote from."""


@dataclass(frozen=True)
class SheetRow:
    row: int  # 1-based row number in the Master tab
    category: str
    product: str
    size: str
    option_1: str | None
    option_2: str | None
    tier_label: str
    qty_from: object
    unit: str
    price: object
    production_time: str
    flags: tuple[str, ...]


# ---------------------------------------------------------------- reading


def read_workbook(path: Path) -> tuple[list[SheetRow], dict[str, str]]:
    if not Path(path).exists():
        raise LoaderError(f"Price file not found: {path}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise LoaderError(f"Can't open {Path(path).name} as an Excel workbook: {exc}") from None
    try:
        if MASTER_TAB not in wb.sheetnames or FLAGS_TAB not in wb.sheetnames:
            raise LoaderError(f"Workbook needs tabs {MASTER_TAB!r} and {FLAGS_TAB!r}")
        rows = _read_master(wb[MASTER_TAB])
        flags = _read_flags(wb[FLAGS_TAB])
    finally:
        wb.close()
    return rows, flags


def _header_index(values: tuple, wanted: dict[str, str]) -> dict[str, int] | None:
    names = [str(v).strip() if v is not None else "" for v in values]
    if not all(title in names for title in wanted.values()):
        return None
    return {key: names.index(title) for key, title in wanted.items()}


def _text(value) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def _read_master(ws) -> list[SheetRow]:
    cols = None
    rows: list[SheetRow] = []
    for row_no, values in enumerate(ws.iter_rows(values_only=True), start=1):
        if cols is None:
            cols = _header_index(values, MASTER_COLUMNS)
            continue
        get = lambda key: values[cols[key]] if cols[key] < len(values) else None  # noqa: E731
        if not _text(get("product")):
            continue
        flag_text = _text(get("flag")) or ""
        rows.append(
            SheetRow(
                row=row_no,
                category=_text(get("category")) or "",
                product=_text(get("product")),
                size=dimslib.normalize_size_label(_text(get("size")) or ""),
                option_1=_text(get("option_1")),
                option_2=_text(get("option_2")),
                tier_label=_text(get("tier_label")) or "",
                qty_from=get("qty_from"),
                unit=_text(get("unit")) or "",
                price=get("price"),
                production_time=_text(get("production_time")) or "Not stated",
                flags=tuple(f.strip() for f in re.split(r"[,;/ ]+", flag_text) if f.strip()),
            )
        )
    if cols is None:
        raise LoaderError(f"{MASTER_TAB}: header row not found (need {', '.join(MASTER_COLUMNS.values())})")
    return rows


def _read_flags(ws) -> dict[str, str]:
    cols = None
    flags: dict[str, str] = {}
    for values in ws.iter_rows(values_only=True):
        if cols is None:
            cols = _header_index(values, {"flag": "Flag", "status": "Status"})
            continue
        flag = _text(values[cols["flag"]]) if cols["flag"] < len(values) else None
        if flag:
            flags[flag] = _text(values[cols["status"]]) or ""
    if cols is None:
        raise LoaderError(f"{FLAGS_TAB}: header row with 'Flag' and 'Status' not found")
    return flags


def read_config(path: Path) -> dict:
    if not Path(path).exists():
        raise LoaderError(f"Config file not found: {path}")
    with open(path, encoding="utf-8") as fh:
        config = yaml.safe_load(fh) or {}
    if not isinstance(config.get("products"), list):
        raise LoaderError("config: 'products' list missing")
    return config


# ---------------------------------------------------------------- building


def slug(text: str | None) -> str:
    if not text:
        return ""
    s = text.lower()
    s = re.sub(r"\s*×\s*", "x", s)
    s = re.sub(r"[^a-z0-9.]+", "-", s)
    return s.strip("-.")


def make_item_id(product_id: str, size: str, option_1: str | None, option_2: str | None) -> str:
    return "/".join(p for p in (product_id, slug(size), slug(option_1), slug(option_2)) if p)


def _decimal(value, what: str, row: int) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise LoaderError(f"{MASTER_TAB} row {row}: {what} is empty (was the file saved without recalculating?)")
    if isinstance(value, bool):
        raise LoaderError(f"{MASTER_TAB} row {row}: {what} {value!r} is not a number")
    try:
        result = Decimal(str(value).strip())
    except InvalidOperation:
        raise LoaderError(f"{MASTER_TAB} row {row}: {what} {value!r} is not a number") from None
    if not result.is_finite() or result < 0:
        raise LoaderError(f"{MASTER_TAB} row {row}: {what} {value!r} is not a valid amount")
    return result


def _qty(value, row: int) -> int:
    q = _decimal(value, "Qty from", row)
    if q != q.to_integral_value() or q < 1:
        raise LoaderError(f"{MASTER_TAB} row {row}: Qty from {value!r} is not a whole number")
    return int(q)


def _product_config(entry: dict, defaults: dict, library: dict) -> dict:
    merged = {**defaults, **entry}
    for key in ("product", "id", "sale_unit", "custom_dims"):
        if not merged.get(key):
            raise LoaderError(f"config: product entry {entry.get('product') or entry.get('id')!r} is missing {key!r}")
    if merged["custom_dims"] not in ("none", "box", "bag", "flat", "area_sqft"):
        raise LoaderError(f"config: {merged['id']}: unknown custom_dims {merged['custom_dims']!r}")
    if merged["below_min_policy"] not in ("bill_at_min", "block"):
        raise LoaderError(f"config: {merged['id']}: below_min_policy must be bill_at_min or block")
    if merged["flagged_price_policy"] not in ("warn", "block"):
        raise LoaderError(f"config: {merged['id']}: flagged_price_policy must be warn or block")
    addons = []
    for addon_id in merged.get("addons") or []:
        spec = library.get(addon_id)
        if not spec:
            raise LoaderError(f"config: {merged['id']}: add-on {addon_id!r} not in addon_library")
        from_sheet = spec.get("price") == "from_sheet"
        addons.append(
            Addon(
                id=addon_id,
                name=spec["name"],
                price=None if from_sheet else Decimal(str(spec["price"])),
                basis=spec["basis"],
                from_sheet=from_sheet,
            )
        )
        if spec["basis"] not in ("per_unit", "per_order"):
            raise LoaderError(f"config: add-on {addon_id!r}: basis must be per_unit or per_order")
    merged["_addons"] = addons
    return merged


def build_catalogue(rows: list[SheetRow], flag_status: dict[str, str], config: dict, source: str = "") -> Catalogue:
    defaults = config.get("defaults") or {}
    library = config.get("addon_library") or {}
    default_notes = list(defaults.get("notes") or [])

    configs: dict[str, dict] = {}
    for entry in config["products"]:
        cfg = _product_config(entry, defaults, library)
        if cfg["product"] in configs:
            raise LoaderError(f"config: product {cfg['product']!r} listed twice")
        configs[cfg["product"]] = cfg
    ids = [c["id"] for c in configs.values()]
    if len(set(ids)) != len(ids):
        raise LoaderError("config: product ids must be unique")

    sheet_products = list(dict.fromkeys(r.product for r in rows))
    missing_in_config = [p for p in sheet_products if p not in configs]
    missing_in_sheet = [p for p in configs if p not in sheet_products]
    if missing_in_config:
        raise LoaderError(f"Products in the sheet but not in config: {missing_in_config}")
    if missing_in_sheet:
        raise LoaderError(f"Products in config but not in the sheet: {missing_in_sheet}")

    products: dict[str, Product] = {}
    for name in sheet_products:
        cfg = configs[name]
        category = next(r.category for r in rows if r.product == name)
        products[cfg["id"]] = Product(
            id=cfg["id"],
            name=name,
            display_name=cfg.get("display_name") or name,
            category=category,
            sale_unit=cfg["sale_unit"],
            yield_factor=int(cfg.get("yield_factor") or 1),
            micro_uom=cfg.get("micro_uom"),
            micro_unit=cfg.get("micro_unit"),
            micro_approx=bool(cfg.get("micro_approx", False)),
            custom_dims=cfg["custom_dims"],
            anchor_group=cfg.get("anchor_group"),
            anchor_match=list(cfg.get("anchor_match") or []),
            below_min_policy=cfg["below_min_policy"],
            max_qty=cfg.get("max_qty"),
            custom_surcharge_pct=Decimal(str(cfg.get("custom_surcharge_pct", 0))),
            custom_round_step=Decimal(str(cfg.get("custom_round_step", "0.50"))),
            gst_rate=Decimal(str(cfg.get("gst_rate", "0.18"))),
            flagged_price_policy=cfg["flagged_price_policy"],
            suggest_more=bool(cfg.get("suggest_more", True)),
            size_label=cfg.get("size_label") or "Size",
            option_labels=dict(cfg.get("option_labels") or {}),
            addons=cfg["_addons"],
            notes=default_notes + list(cfg.get("notes") or []),
        )
    by_name = {p.name: p for p in products.values()}

    # Group rows into items; sample rows become the item's sample charge.
    items: dict[tuple, Item] = {}
    samples: dict[tuple, Decimal] = {}
    sample_rows = tier_rows = 0
    for r in rows:
        key = (r.product, r.size, r.option_1, r.option_2)
        price = _decimal(r.price, "Price", r.row)
        if r.tier_label.strip().lower() == SAMPLE_TIER:
            samples[key] = price
            sample_rows += 1
            continue
        tier_rows += 1
        product = by_name[r.product]
        item = items.get(key)
        if item is None:
            item = items[key] = Item(
                id=make_item_id(product.id, r.size, r.option_1, r.option_2),
                product_id=product.id,
                size=r.size,
                option_1=r.option_1,
                option_2=r.option_2,
                tiers=[],
                production_time=r.production_time,
                unit=r.unit,
            )
        qty_from = _qty(r.qty_from, r.row)
        if any(t.qty_from == qty_from for t in item.tiers):
            raise LoaderError(
                f"{MASTER_TAB} row {r.row}: repeated Qty from {qty_from} for {r.product} / {r.size}"
                f" / {r.option_1 or '-'} / {r.option_2 or '-'}"
            )
        item.tiers.append(Tier(qty_from=qty_from, label=r.tier_label, price=price, flags=r.flags, row=r.row))

    warnings: list[str] = []
    by_id: dict[str, Item] = {}
    for key, item in items.items():
        item.tiers.sort(key=lambda t: t.qty_from)
        item.sample_charge = samples.pop(key, None)
        if item.id in by_id:
            raise LoaderError(f"Two items share the id {item.id!r}; make their Size / Spec or options differ")
        by_id[item.id] = item
        product = products[item.product_id]
        if product.supports_custom:
            parsed = dimslib.parse_size(item.size)
            if len(parsed) != len(dimslib.DIM_NAMES[product.custom_dims]) or not all(parsed):
                raise LoaderError(
                    f"{MASTER_TAB} row {item.tiers[0].row}: can't read {product.custom_dims} size from {item.size!r}"
                )
            item.dims = parsed
            item.metric = dimslib.metric(product.custom_dims, parsed)
        prices = [t.price for t in item.tiers]
        if any(b >= a for a, b in zip(prices, prices[1:])):
            msg = f"Price doesn't fall as quantity rises: {item.id} ({' / '.join(str(p) for p in prices)})"
            warnings.append(msg)
            log.warning(msg)
        product.items.append(item)
    if samples:
        orphan = next(iter(samples))
        raise LoaderError(f"Sample cost row for {orphan} has no priced tiers")

    # Anchor-group members must share tier breakpoints.
    groups: dict[str, list[Item]] = {}
    for product in products.values():
        if product.supports_custom:
            groups.setdefault(product.anchor_group or product.id, []).extend(product.items)
    for group, members in groups.items():
        breakpoints = {tuple(i.breakpoints) for i in members}
        if len(breakpoints) > 1:
            raise LoaderError(f"Anchor group {group!r} mixes tier breakpoints {sorted(breakpoints)}")

    return Catalogue(
        products=products,
        items=by_id,
        categories=list(dict.fromkeys(p.category for p in products.values())),
        flag_status=flag_status,
        sample_rows=sample_rows,
        tier_rows=tier_rows,
        warnings=warnings,
        show_invoice_billing=bool(defaults.get("show_invoice_billing", False)),
        loaded_at=datetime.now(timezone.utc),
        source_file=source,
    )


def load(data_file: Path, config_file: Path) -> Catalogue:
    rows, flags = read_workbook(data_file)
    config = read_config(config_file)
    catalogue = build_catalogue(rows, flags, config, source=Path(data_file).name)
    log.info(
        "Loaded %s: %d tier rows, %d items, %d products, %d open flags",
        Path(data_file).name,
        catalogue.tier_rows,
        len(catalogue.items),
        len(catalogue.products),
        catalogue.open_flag_count,
    )
    return catalogue
