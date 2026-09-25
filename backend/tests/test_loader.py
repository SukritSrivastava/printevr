"""L1-L4 and loader rules (BRD 4.4)."""
import copy
import dataclasses

import openpyxl
import pytest

from app.dims import metric, normalize_size_label, parse_size
from app.loader import LoaderError, load, make_item_id


def test_l1_counts(catalogue):
    # The BRD's 284 counts each mailer-bag tier as its own item because the sheet's size
    # text repeats "in" ("6 × 8 in in in"). Normalised, the sheet holds 233 items.
    assert catalogue.tier_rows == 950
    assert len(catalogue.items) == 233
    assert len(catalogue.products) == 27
    assert len(catalogue.categories) == 14
    assert catalogue.sample_rows == 12


def test_l1_every_row_is_accounted_for(catalogue):
    assert sum(len(i.tiers) for i in catalogue.items.values()) == 950
    assert sum(1 for i in catalogue.items.values() if i.sample_charge is not None) == 12


def test_mailer_sizes_are_normalised(catalogue):
    courier = catalogue.products["courier_bags"]
    assert [i.size for i in courier.items] == [
        "6 × 8 in", "8 × 10 in", "10 × 14 in", "12 × 16 in", "14 × 18 in", "16 × 22 in"
    ]
    assert all(i.breakpoints == [200, 500, 1000, 2500, 5000] for i in courier.items)


def test_l2_blank_price_stops_startup_and_names_the_row(sheet, rebuild):
    rows = list(sheet[0])
    victim = rows[100]
    rows[100] = dataclasses.replace(victim, price=None)
    with pytest.raises(LoaderError, match=f"row {victim.row}"):
        rebuild(rows=rows)


def test_l2_blank_price_in_file(settings, tmp_path):
    wb = openpyxl.load_workbook(settings.data_file, data_only=True)
    wb["Master Price List"].cell(row=50, column=10).value = None
    path = tmp_path / "blank.xlsx"
    wb.save(path)
    with pytest.raises(LoaderError, match="row 50"):
        load(path, settings.config_file)


def test_l2_file_saved_without_recalculating_is_refused(settings, tmp_path):
    wb = openpyxl.load_workbook(settings.data_file)  # keeps formulas; cached values are lost on save
    path = tmp_path / "uncalculated.xlsx"
    wb.save(path)
    with pytest.raises(LoaderError, match="empty"):
        load(path, settings.config_file)


def test_l3_duplicated_tier_row_stops_startup(sheet, rebuild):
    rows = list(sheet[0])
    rows.insert(50, rows[50])
    with pytest.raises(LoaderError, match="repeated Qty from"):
        rebuild(rows=rows)


def test_l4_price_does_not_fall_warning_on_exactly_11_items(catalogue):
    assert len(catalogue.warnings) == 11
    flagged = " ".join(catalogue.warnings)
    for item in (
        "rigid_boxes/8x10x2.5-in/magnetic-slider",
        "rigid_boxes/18x12x4-in/magnetic-slider",
        "vc_textured/300-gsm/no-lamination/double-side",
        "id_cards/holder",
        "frosted_bags/10x14-in",
    ):
        assert item in flagged


def test_product_missing_from_config_stops_startup(sheet, rebuild):
    config = copy.deepcopy(sheet[2])
    config["products"] = [p for p in config["products"] if p["id"] != "rigid_boxes"]
    with pytest.raises(LoaderError, match="not in config"):
        rebuild(config=config)


def test_config_product_missing_from_sheet_stops_startup(sheet, rebuild):
    config = copy.deepcopy(sheet[2])
    config["products"].append({"product": "Gift Hampers", "id": "hampers", "sale_unit": "pc", "custom_dims": "none"})
    with pytest.raises(LoaderError, match="not in the sheet"):
        rebuild(config=config)


def test_anchor_group_with_different_breakpoints_stops_startup(sheet, rebuild):
    rows = [
        dataclasses.replace(r, qty_from=150) if r.product == "Small Carry Bags" and r.qty_from == 100 else r
        for r in sheet[0]
    ]
    with pytest.raises(LoaderError, match="carry_bags"):
        rebuild(rows=rows)


def test_unparseable_custom_size_stops_startup(sheet, rebuild):
    rows = [
        dataclasses.replace(r, size="Assorted") if r.product == "Rigid Boxes" and r.size == "3 × 3 × 2 in" else r
        for r in sheet[0]
    ]
    with pytest.raises(LoaderError, match="can't read box size"):
        rebuild(rows=rows)


def test_flag_column_with_several_ids(sheet, rebuild):
    rows = list(sheet[0])
    i = next(n for n, r in enumerate(rows) if r.product == "Rigid Boxes" and r.size == "3 × 3 × 2 in")
    rows[i] = dataclasses.replace(rows[i], flags=("F1", "F5"))
    cat = rebuild(rows=rows)
    assert cat.item("rigid_boxes/3x3x2-in/top-bottom").tiers[0].flags == ("F1", "F5")


def test_size_parsing():
    assert parse_size("3 × 3 × 2 in") == (3, 3, 2)
    assert parse_size("7 × 9 × 3 in (H × L × S)") == (7, 9, 3)
    assert parse_size("8 × 8 × 1.5 in (listed twice)") == (8, 8, 1.5)
    assert parse_size("29x9") == (29, 9)
    assert parse_size("4 X 5 * 2.5") == (4, 5, 2.5)
    assert normalize_size_label("6 × 8 in in in in") == "6 × 8 in"


def test_metrics():
    assert metric("box", parse_size("3.5 × 3.5 × 2")) == 52.5
    assert metric("bag", parse_size("10 × 10 × 3")) == 290
    assert metric("flat", parse_size("29 × 9")) == 261


def test_item_ids():
    assert make_item_id("rigid_boxes", "3 × 3 × 2 in", "Top-Bottom", None) == "rigid_boxes/3x3x2-in/top-bottom"
    assert (
        make_item_id("vc_standard", "300 GSM", "Gloss / Matte", "Double side")
        == "vc_standard/300-gsm/gloss-matte/double-side"
    )


def test_sample_rows_become_sample_charges(catalogue):
    assert str(catalogue.item("sticker_sheets/paper").sample_charge) == "250"
    assert str(catalogue.item("paper_printing/90-120-gsm").sample_charge) == "200"
